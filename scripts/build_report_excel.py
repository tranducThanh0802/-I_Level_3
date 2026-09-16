#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Xuất báo cáo đánh giá (§6) ra Excel: report.xlsx (nhiều sheet).
Nguồn: agent-team/logs/runs.jsonl + memory/bugs/. Quy tắc §6: median, tách chờ ngoài.
Chạy: python3 scripts/build_report_excel.py  rồi mở report.xlsx
"""
import os, re, json, glob, statistics, datetime
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AT = os.path.join(ROOT, "agent-team")
OUT = os.path.join(ROOT, "report.xlsx")
try:
    PROJECT = json.load(open(os.path.join(ROOT, "project.config.json"), encoding="utf-8")).get("name", "Agent Team")
except Exception:
    PROJECT = "Agent Team"

HEAD = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="2F6FEB")
TITLE = Font(bold=True, size=14)
WARN = Font(bold=True, color="C00000")
THIN = Border(*[Side(style="thin", color="DDDDDD")] * 4)


def med(xs):
    xs = [x for x in xs if x is not None]
    return round(statistics.median(xs), 1) if xs else None


def dur_min(r):
    try:
        s = datetime.datetime.fromisoformat(r["started_at"]); e = datetime.datetime.fromisoformat(r["ended_at"])
        return round(max((e - s).total_seconds() / 60 - r.get("wait_external_ms", 0) / 60000, 0), 1)
    except Exception:
        return None


def saving_pct(r):
    b = r.get("human_baseline_min"); d = dur_min(r)
    return round((b - d) / b * 100) if (b and d is not None and b > 0) else None


def fm_of(text):
    fm = {}
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            for line in text[3:end].splitlines():
                m = re.match(r"^([A-Za-z_]+):\s*(.*)$", line)
                if m:
                    fm[m.group(1)] = m.group(2).strip()
    return fm


# --- load ---
runs = []
rp = os.path.join(AT, "logs", "runs.jsonl")
if os.path.exists(rp):
    for l in open(rp, encoding="utf-8"):
        l = l.strip()
        if l:
            try:
                runs.append(json.loads(l))
            except Exception:
                pass
bugs = [fm_of(open(f, encoding="utf-8").read()) for f in glob.glob(os.path.join(AT, "memory", "bugs", "*.md"))
        if os.path.basename(f).upper() != "README.MD"]

total = len(runs)
used = sum(1 for r in runs if r.get("outcome") == "used")
dropped = sum(1 for r in runs if r.get("outcome") == "discarded")
n_demo = sum(1 for r in runs if r.get("demo") is True)
fix_types = [mf.get("type", "?") for r in runs for mf in r.get("manual_fixes", [])]
test_runs = [r for r in runs if r.get("agent") == "test"]
tgreen = sum(1 for r in test_runs if any(s.get("result") == "green" for s in r.get("steps", [])))
dedup_saved = sum(max(int(b.get("count", 1) or 1) - 1, 0) for b in bugs if str(b.get("count", "1")).isdigit())

wb = Workbook()


def style_header(ws, row=1):
    for c in ws[row]:
        c.font = HEAD; c.fill = HFILL; c.alignment = Alignment(vertical="center")


def autofit(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# Sheet 1: Tổng quan
ws = wb.active; ws.title = "Tổng quan"
ws["A1"] = f"{PROJECT} — Báo cáo đánh giá (§6)"; ws["A1"].font = TITLE
ws["A2"] = f"Xuất: {datetime.date.today().isoformat()} · median (không mean) · chờ ngoài tách riêng"
ws["A2"].font = Font(italic=True, color="808080")
row = 4
if n_demo:
    ws[f"A{row}"] = f"⚠️ DỮ LIỆU DEMO: {n_demo}/{total} dòng — KHÔNG phải số thật (xoá: seed_demo_data.py --clear)"
    ws[f"A{row}"].font = WARN; row += 2
kpis = [
    ("Chỉ số", "Giá trị"),
    ("Tổng lần chạy (có log)", total),
    ("Tỷ lệ dùng ngay", f"{round(used/total*100) if total else 0}%"),
    ("Tỷ lệ bỏ giữa chừng", f"{round(dropped/total*100) if total else 0}%"),
    ("Tiết kiệm vs làm tay (median, mục tiêu ≥40%)",
     (lambda s: f"{s}%" if s is not None else "— (thiếu baseline)")(med([saving_pct(r) for r in runs]))),
    ("Tổng lần sửa tay", len(fix_types)),
    ("Test XANH / ĐỎ", f"{tgreen} / {len(test_runs)-tgreen}"),
    ("Bug gộp trùng (đỡ mở bản ghi mới)", dedup_saved),
    ("Số bug đang theo dõi", len(bugs)),
]
for i, (a, b) in enumerate(kpis):
    ws[f"A{row+i}"] = a; ws[f"B{row+i}"] = b
    if i == 0:
        ws[f"A{row+i}"].font = HEAD; ws[f"A{row+i}"].fill = HFILL
        ws[f"B{row+i}"].font = HEAD; ws[f"B{row+i}"].fill = HFILL
r2 = row + len(kpis) + 1
ws[f"A{r2}"] = "Ngưỡng ĐẠT (§6):"; ws[f"A{r2}"].font = Font(bold=True)
for i, t in enumerate(["Bậc ≥2", "Giảm ≥40% thời gian việc lặp", "Bỏ dở <20%", "Bug quay lại <10%",
                       "So sánh: cùng loại, ≥5 mẫu/bên, dùng median"]):
    ws[f"A{r2+1+i}"] = "• " + t
autofit(ws, [42, 22])

# Sheet 2: Thời gian theo loại
ws2 = wb.create_sheet("Thời gian theo loại")
ws2.append(["Luồng", "Median thời gian làm (phút)", "Median chờ ngoài (phút)",
            "Tiết kiệm vs làm tay (median)", "Số mẫu", "Đủ ≥5?"])
byf = {}
for r in runs:
    byf.setdefault(r.get("flow", "?"), []).append(r)
for flow, rs in sorted(byf.items()):
    sv = med([saving_pct(r) for r in rs])
    ws2.append([flow, med([dur_min(r) for r in rs]),
                med([r.get("wait_external_ms", 0) / 60000 for r in rs]),
                (sv / 100 if sv is not None else None), len(rs),
                "Đủ" if len(rs) >= 5 else "CHƯA (cần ≥5)"])
for r in range(2, ws2.max_row + 1):
    ws2.cell(row=r, column=4).number_format = "0%"
style_header(ws2); autofit(ws2, [16, 26, 24, 28, 10, 16])

# Sheet 3: Chi tiết lần chạy
ws3 = wb.create_sheet("Chi tiết lần chạy")
cols = ["run_id", "agent", "flow", "task", "Size", "Phút làm", "Baseline tay", "Tiết kiệm",
        "Phút chờ ngoài", "outcome", "Test", "Sửa tay (loại)", "Bàn giao đủ?", "demo?"]
ws3.append(cols)
for r in runs:
    tr = next((s.get("result") for s in r.get("steps", []) if s.get("step", "").startswith("chạy")), "")
    ho = r.get("handoff", {})
    sv = saving_pct(r)
    ws3.append([r.get("run_id", ""), r.get("agent", ""), r.get("flow", ""), r.get("task", ""),
                r.get("size", ""), dur_min(r), r.get("human_baseline_min", ""),
                (sv / 100 if sv is not None else ""),
                round(r.get("wait_external_ms", 0) / 60000, 1), r.get("outcome", ""),
                tr, ", ".join(mf.get("type", "") for mf in r.get("manual_fixes", [])),
                "có" if all([ho.get("code"), ho.get("video"), ho.get("before_after")]) else "thiếu",
                "demo" if r.get("demo") else ""])
for r in range(2, ws3.max_row + 1):
    ws3.cell(row=r, column=8).number_format = "0%"
style_header(ws3); autofit(ws3, [34, 10, 10, 26, 6, 9, 12, 10, 14, 12, 8, 26, 12, 8])
ws3.freeze_panes = "A2"

# Sheet 4: Sửa tay (top loại)
ws4 = wb.create_sheet("Sửa tay")
ws4.append(["Loại sửa tay", "Số lần"])
for t, c in Counter(fix_types).most_common():
    ws4.append([t, c])
if not fix_types:
    ws4.append(["(chưa có)", 0])
style_header(ws4); autofit(ws4, [30, 12])

# Sheet 5: Bug
ws5 = wb.create_sheet("Bug")
ws5.append(["ID", "key", "severity", "status", "count", "assignee", "nguồn"])
for b in bugs:
    ws5.append([b.get("id", ""), b.get("key", ""), b.get("severity", ""), b.get("status", ""),
                b.get("count", ""), b.get("assignee", ""), b.get("source", "")])
style_header(ws5); autofit(ws5, [10, 26, 10, 14, 8, 16, 10])

wb.save(OUT)
print(f"Đã tạo: {OUT} ({total} lần chạy, {len(bugs)} bug)")
